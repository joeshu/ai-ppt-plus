#!/usr/bin/env python3
from __future__ import annotations
import argparse
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

W,H=13.333333,7.5
FONT='Noto Sans CJK SC'
FIXED_DT=datetime(2026,9,11,0,0,0,tzinfo=timezone.utc)
ZIP_DT=(2026,9,11,0,0,0)

def rgb(s): return RGBColor.from_string(s)
def stabilize(prs):
    cp=prs.core_properties; cp.title='China Unicom goals replay fixture'; cp.subject='Deterministic fixed-reference reconstruction replay'; cp.author='OpenAI'; cp.last_modified_by='OpenAI'; cp.created=FIXED_DT.replace(tzinfo=None); cp.modified=FIXED_DT.replace(tzinfo=None); cp.revision=1

def canonicalize(path:Path):
    tmp=path.with_suffix(path.suffix+'.tmp')
    with ZipFile(path,'r') as zin, ZipFile(tmp,'w',compression=ZIP_DEFLATED,compresslevel=9) as zout:
        for name in sorted(zin.namelist()):
            info=ZipInfo(name,ZIP_DT); info.compress_type=ZIP_DEFLATED; info.external_attr=0o600<<16; zout.writestr(info,zin.read(name))
    tmp.replace(path)

def add_text(slide,text,x,y,w,h,size=14,color='111111',bold=False,align='left'):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=text; p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(align,PP_ALIGN.LEFT)
    for r in p.runs: r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return box

def shape(slide,kind,x,y,w,h,fill='FFFFFF',line=None):
    shp=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); shp.line.color.rgb=rgb(line or fill); return shp

def crop_assets(ref:Path,out:Path):
    out.mkdir(parents=True,exist_ok=True); im=Image.open(ref).convert('RGB'); sx=im.width/1536; sy=im.height/1024
    specs={'rural':(46,291,132,132),'community':(640,291,132,132),'target':(1218,304,92,100),'goal1':(130,732,96,96),'goal2':(433,732,96,96),'goal3':(720,732,96,96),'goal4':(1009,732,96,96),'goal5':(1287,732,96,96),'logo':(1280,8,230,95),'bottom-wave':(0,944,1536,80)}
    result={}
    for name,(x,y,w,h) in specs.items():
        box=(round(x*sx),round(y*sy),round((x+w)*sx),round((y+h)*sy)); p=out/f'{name}.jpg'; im.crop(box).save(p,quality=90); result[name]=p
    return result

def add_pic(slide,p,x,y,w,h): slide.shapes.add_picture(str(p),Inches(x),Inches(y),Inches(w),Inches(h))
def add_tag(slide,n,title,y):
    shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,0.22,y,0.38,0.34,'C40000'); add_text(slide,n,0.31,y+0.01,0.19,0.28,18,'FFFFFF',True,'center'); add_text(slide,title,0.69,y-0.01,2.3,0.36,18,'C40000',True)

def build(reference:Path,outdir:Path):
    outdir.mkdir(parents=True,exist_ok=True); a=crop_assets(reference,outdir/'assets')
    prs=Presentation(); stabilize(prs); prs.slide_width=Inches(W); prs.slide_height=Inches(H); s=prs.slides.add_slide(prs.slide_layouts[6]); add_pic(s,reference,0,0,W,H); prs.save(outdir/'source-reference.pptx'); canonicalize(outdir/'source-reference.pptx')
    prs=Presentation(); stabilize(prs); prs.slide_width=Inches(W); prs.slide_height=Inches(H); s=prs.slides.add_slide(prs.slide_layouts[6])
    shape(s,MSO_SHAPE.RECTANGLE,0,0,1.43,0.8,'C90000'); add_text(s,'夺高点\n争上场',0.28,0.12,0.9,0.52,18,'FFF500',True); add_text(s,'第二部分',1.72,0.22,1.34,0.38,22,'5A5A5A',True); add_text(s,'目标（干多少）',3.29,0.16,3.45,0.48,28,'C40000',True)
    shape(s,MSO_SHAPE.RECTANGLE,1.43,0.79,11.9,0.02,'C40000'); add_pic(s,a['logo'],11.2,0.10,1.8,0.55); add_text(s,'紧盯标杆打造目标，以',1.8,0.96,3.0,0.35,18,'111111',True); add_text(s,'移宽双新融合',4.6,0.96,1.95,0.35,19,'E00000',True); add_text(s,'为主攻方向，全力完成三季度发展任务！',6.55,0.96,4.9,0.35,18,'111111',True)
    add_tag(s,'1','标杆打造目标',1.47); shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,0.22,1.93,12.9,1.45,'FFFFFF','EF8A8A'); add_pic(s,a['rural'],0.44,2.13,1.0,0.92); shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,1.34,2.36,1.12,0.43,'CC0000'); add_text(s,'农村市场',1.48,2.39,0.86,0.30,16,'FFFFFF',True,'center'); shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,2.62,2.13,2.55,0.96,'FFF5F5'); add_text(s,'50个',2.76,2.28,0.66,0.33,23,'D90000',True); add_text(s,'重点网格',3.37,2.31,0.82,0.30,14,'111111',True); add_text(s,'30户',4.26,2.28,0.73,0.33,23,'D90000',True); add_text(s,'其余网格',2.76,2.67,0.85,0.28,14,'111111',True); add_text(s,'20户',3.67,2.63,0.74,0.32,22,'D90000',True)
    shape(s,MSO_SHAPE.RECTANGLE,5.42,2.12,0.01,1.05,'F2A3A3'); add_pic(s,a['community'],5.63,2.13,1.0,0.92); shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,6.52,2.36,1.11,0.43,'CC0000'); add_text(s,'社区市场',6.63,2.39,0.9,0.30,16,'FFFFFF',True,'center'); shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,7.76,2.13,2.27,0.96,'FFF5F5'); add_text(s,'每个标杆社区',7.9,2.53,1.26,0.28,14,'111111',True); add_text(s,'20户',9.16,2.47,0.73,0.37,22,'D90000',True)
    shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,10.32,1.93,2.80,1.45,'FFFDF2'); add_pic(s,a['target'],10.55,2.28,0.76,0.68); add_text(s,'目标要求',11.34,2.20,1.2,0.34,18,'C40000',True); add_text(s,'全部以79元及以上',11.34,2.61,1.48,0.26,13,'111111',True); add_text(s,'移宽双新融合为目标',11.34,2.89,1.56,0.26,13,'111111',True)
    add_tag(s,'2','推进节奏安排',3.57); shape(s,MSO_SHAPE.RECTANGLE,0.8,4.77,11.3,0.04,'F22B35'); shape(s,MSO_SHAPE.CHEVRON,11.92,4.66,0.38,0.26,'E9101D')
    for x,w,t in [(1.28,0.68,'6.28'),(4.30,1.20,'7.1 – 7.20'),(8.05,1.20,'7.21 – 7.31'),(10.98,0.78,'8.3')]: shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,4.02,w,0.23,'D00000'); add_text(s,t,x,4.02,w,0.23,11,'FFFFFF',True,'center')
    add_text(s,'上报目标',1.24,4.28,0.90,0.28,14,'111111',True,'center'); shape(s,MSO_SHAPE.OVAL,2.18,4.21,0.54,0.54,'F25E65'); add_text(s,'启动',2.24,4.31,0.42,0.25,12,'FFFFFF',True,'center'); shape(s,MSO_SHAPE.OVAL,2.39,4.72,0.12,0.12,'F25E65'); add_text(s,'一阶段：标杆打造',4.07,4.29,1.65,0.26,13,'C40000',True,'center'); add_text(s,'高新完成4个社区，示范完成3个社区+1个行政村',3.30,4.55,3.1,0.24,10,'111111',False,'center'); add_text(s,'二阶段：培优补差',7.82,4.29,1.72,0.26,13,'C40000',True,'center'); add_text(s,'一阶段已达标的行政村和社区\n二阶段再完成10户目标',7.62,4.53,2.14,0.40,10,'111111',False,'center'); add_text(s,'月度复盘表彰大会',10.63,4.29,1.55,0.27,13,'111111',True,'center'); shape(s,MSO_SHAPE.OVAL,12.20,4.39,0.64,0.64,'E9111C'); add_text(s,'工作\n结束',12.27,4.50,0.50,0.38,12,'FFFFFF',True,'center')
    add_tag(s,'3','目标要求',5.02); shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,0.25,5.35,12.7,1.65,'FFFFFF','D9D9D9')
    goals=[(1.20,'goal1','对标先进','C40000','对标前期先进网格找差距\n明确提升方向和路径'),(3.83,'goal2','聚焦双新','F04A20','坚持以移宽双新融合为主线\n提升发展效率和质量'),(6.32,'goal3','全员参与','E9A900','网格CEO带头、全员上阵\n形成合力，全面推进'),(8.82,'goal4','以质取胜','1565C0','注重发展质量和客户体验\n打造可复制的标杆样板'),(11.28,'goal5','持续增长','3E7C3F','以标杆带动整体提升\n确保三季度目标达成')]
    for i,(x,img,title,color,body) in enumerate(goals):
        add_pic(s,a[img],x-0.12,5.52,0.72,0.72); add_text(s,title,x-0.32,6.25,1.15,0.28,16,color,True,'center'); add_text(s,body,x-0.62,6.58,1.75,0.48,10,'111111',False,'center')
        if i<4: shape(s,MSO_SHAPE.RECTANGLE,x+1.55,5.55,0.01,1.34,'C8C8C8')
    add_pic(s,a['bottom-wave'],0,7.02,W,0.48); prs.save(outdir/'editable.pptx'); canonicalize(outdir/'editable.pptx')

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--reference',required=True,type=Path); ap.add_argument('--output-dir',required=True,type=Path); args=ap.parse_args(); build(args.reference,args.output_dir)
if __name__=='__main__': main()
