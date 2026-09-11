#!/usr/bin/env python3
"""Real checked-in replay for the China Unicom goals 16:9 reconstruction."""
from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from zipfile import is_zipfile

import numpy as np
from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
CASE_DIR = ROOT / "evals" / "case-replay-unicom-goals-20260911"
CASE = CASE_DIR / "case.json"
MATERIALIZER = ROOT / "scripts" / "materialize_replay_artifact.py"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def materialize(manifest: Path, output: Path) -> None:
    if manifest.suffix.lower() in {".pptx", ".png", ".jpg", ".jpeg"}:
        shutil.copy2(manifest, output)
        return
    completed = subprocess.run(
        [sys.executable, str(MATERIALIZER), str(manifest), "--output", str(output)],
        cwd=ROOT, capture_output=True, text=True, check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr


def render(deck: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftoppm = shutil.which("pdftoppm")
    assert soffice and pdftoppm, "LibreOffice/Poppler renderer is required for real replay"
    output_dir.mkdir(parents=True, exist_ok=True)
    profile = output_dir / "lo-profile"
    profile.mkdir()
    converted = subprocess.run(
        [soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(deck)],
        capture_output=True, text=True, check=False,
    )
    assert converted.returncode == 0, converted.stdout + converted.stderr
    pdf = output_dir / f"{deck.stem}.pdf"
    assert pdf.is_file(), converted.stdout + converted.stderr
    converted = subprocess.run(
        [pdftoppm, "-singlefile", "-png", "-r", "144", str(pdf), str(output_dir / "render")],
        capture_output=True, text=True, check=False,
    )
    assert converted.returncode == 0, converted.stdout + converted.stderr
    rendered = output_dir / "render.png"
    assert rendered.is_file()
    return rendered


def ssim_gray(a: np.ndarray, b: np.ndarray) -> float:
    a = a.astype(np.float64)
    b = b.astype(np.float64)
    c1 = (0.01 * 255) ** 2
    c2 = (0.03 * 255) ** 2
    ma, mb = a.mean(), b.mean()
    va, vb = a.var(), b.var()
    cov = ((a - ma) * (b - mb)).mean()
    return float(((2 * ma * mb + c1) * (2 * cov + c2)) / ((ma * ma + mb * mb + c1) * (va + vb + c2)))


def blurred_layout_similarity(rendered: Path, reference: Path) -> float:
    with Image.open(rendered) as rim, Image.open(reference) as sim:
        target = rim.convert("L")
        source = sim.convert("L").resize(target.size, Image.Resampling.LANCZOS)
        radius = max(2.0, min(target.size) / 110.0)
        target = target.filter(ImageFilter.GaussianBlur(radius))
        source = source.filter(ImageFilter.GaussianBlur(radius))
        return ssim_gray(np.asarray(target), np.asarray(source))


def main() -> int:
    spec = json.loads(CASE.read_text(encoding="utf-8"))
    replay = spec["replay"]
    expected = spec["expected"]
    with tempfile.TemporaryDirectory(prefix="unicom-goals-replay-") as temporary:
        work = Path(temporary)
        source = work / "source-reference.pptx"
        candidate = work / "editable.pptx"
        reference = work / "reference.png"
        materialize(ROOT / replay["source_deck"], source)
        materialize(ROOT / replay["candidate_deck"], candidate)
        materialize(ROOT / replay["reference_image"], reference)
        assert is_zipfile(source) and is_zipfile(candidate)
        assert digest(source) == spec["approved_artifact_hashes"]["source_deck"]
        assert digest(candidate) == spec["approved_artifact_hashes"]["candidate_deck"]
        assert digest(reference) == spec["approved_artifact_hashes"]["reference_image"]

        prs = Presentation(str(candidate))
        assert len(prs.slides) == 1
        assert abs((prs.slide_width / prs.slide_height) - (16 / 9)) < 0.01
        flat = list(iter_shapes(prs.slides[0].shapes))
        assert sum(1 for shape in flat if getattr(shape, "has_table", False)) == expected["native_tables"] == 0
        assert not [shape for shape in flat if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width >= prs.slide_width * 0.95 and shape.height >= prs.slide_height * 0.95]
        assert len([shape for shape in flat if getattr(shape, "has_text_frame", False) and shape.text.strip()]) >= expected["min_native_text_objects"]
        assert len(flat) >= expected["min_total_objects"]

        rendered = render(candidate, work / "render")
        similarity = blurred_layout_similarity(rendered, reference)
        assert similarity >= expected["replay_blurred_similarity_min"], similarity

        mutated = work / "mutated.pptx"
        shutil.copy2(candidate, mutated)
        deck = Presentation(str(mutated))
        flat_mut = list(iter_shapes(deck.slides[0].shapes))
        text = next(shape for shape in flat_mut if getattr(shape, "has_text_frame", False) and "目标" in shape.text)
        old_text = text.text
        text.text = old_text + "·M"
        movable = next(shape for shape in flat_mut if shape is not text and shape.width > 0 and shape.height > 0)
        old_left = movable.left
        movable.left = old_left + 12700
        deck.save(mutated)
        reopened = Presentation(str(mutated))
        flat_after = list(iter_shapes(reopened.slides[0].shapes))
        assert any(getattr(shape, "has_text_frame", False) and shape.text == old_text + "·M" for shape in flat_after)
        assert any(shape.left == old_left + 12700 for shape in flat_after)

    print(f"unicom-goals-16x9-fixed-reference-01 real replay: ok (blurred_similarity={similarity:.6f})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
