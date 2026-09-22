from __future__ import annotations

import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches
from pptx_primitives import _set_run_fonts
from runtime_fonts import runtime_font_evidence


def test_runtime_font_resolver_never_accepts_latin_fallback_as_cjk():
    evidence = runtime_font_evidence("Noto Sans CJK SC")
    assert evidence["repository_font_binary_required"] is False
    assert evidence["files"]
    assert evidence["valid"] == any(item["cjk_coverage_valid"] for item in evidence["coverage"])
    for item in evidence["coverage"]:
        if Path(item["path"]).name.startswith("DejaVuSans"):
            assert item["cjk_coverage_valid"] is False


def test_cjk_run_binds_east_asian_and_complex_script_typefaces():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "中文 English 123"
    _set_run_fonts(run, "Noto Sans CJK SC")
    rpr = run._r.get_or_add_rPr()
    assert run.font.name == "Noto Sans CJK SC"
    assert rpr.find(qn("a:latin")).get("typeface") == "Noto Sans CJK SC"
    assert rpr.find(qn("a:ea")).get("typeface") == "Noto Sans CJK SC"
    assert rpr.find(qn("a:cs")).get("typeface") == "Noto Sans CJK SC"


def test_repository_does_not_track_font_binaries():
    repo = ROOT.parent
    proc = subprocess.run(["git", "ls-files", "assets/fonts", "ai-ppt-editable/assets/fonts"], cwd=repo, capture_output=True, text=True, check=True)
    tracked = [line for line in proc.stdout.splitlines() if Path(line).suffix.lower() in {".ttf", ".otf", ".ttc"}]
    assert tracked == []
