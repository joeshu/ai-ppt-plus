"""Real PowerPoint smoke test; executed only by the labelled Windows runner."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "powerpoint-smoke-output"
OUT.mkdir(exist_ok=True)
deck_path = OUT / "fixture.pptx"
asset_path = OUT / "transparent.png"

image = Image.new("RGBA", (160, 160), (0, 0, 0, 0))
ImageDraw.Draw(image).ellipse((16, 16, 144, 144), fill=(220, 0, 30, 220))
image.save(asset_path)

deck = Presentation()
deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
for index in range(2):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.2))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = f"中文换行与 PowerPoint 渲染 {index + 1} — 100% / 25℃"
    run.font.name, run.font.size, run.font.bold = "Microsoft YaHei", Pt(28), True
    slide.shapes.add_picture(str(asset_path), Inches(1), Inches(3), width=Inches(1.4), height=Inches(1.4))
deck.save(deck_path)

report = OUT / "render-report.json"
result = subprocess.run([sys.executable, str(ROOT / "ai-ppt-editable/scripts/render_powerpoint.py"), str(deck_path), "--output-dir", str(OUT / "render"), "--report", str(report)], check=False)
payload = json.loads(report.read_text(encoding="utf-8")) if report.is_file() else {}
assert result.returncode == 0, payload
assert [Path(path).name for path in payload.get("pages", [])] == ["slide-1.png", "slide-2.png"], payload
assert all(Image.open(path).size == (1920, 1080) for path in payload["pages"])
print(json.dumps({"ok": True, "pages": payload["pages"]}, ensure_ascii=False))
