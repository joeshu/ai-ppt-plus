from __future__ import annotations

import argparse
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from asset_placement import alpha_centroid_fit
from audit_pptx_layers import audit
from patch_chart_blank_series import patch_chart
from ppt_text_fit import _tokens, best_fit
from slice_grid import _detect_grid, _square_repack
from validate_transparent_assets import validate


def _rgba(path: Path, offset: int = 12) -> None:
    image = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
    ImageDraw.Draw(image).ellipse((offset, 14, offset + 28, 48), fill=(220, 0, 0, 255))
    image.save(path)


def test_text_fit_reports_reference_box_deficit_and_uses_bundled_font():
    args = argparse.Namespace(
        text="看 diff、看运行结果", box="180x34", font="Microsoft YaHei", font_file=None,
        bold=False, min_pt=8.0, max_pt=26.0, max_lines=1, line_spacing=1.06,
        width_safety=0.92, height_safety=0.95, render_fudge=1.01,
        target_pt=24.0, slide_px="1672x941", slide_in="13.333333x7.505",
    )
    result = best_fit(args)
    assert result["recommended_pt"] >= 8
    assert result["target"]["required_box_px"][0] > 0
    assert Path(result["font_path"]).is_file()
    assert result["reference_scale"] <= 1.1
    assert result["target_pt"] == 24.0
    assert result["required_box_px"] == result["target"]["required_box_px"]
    assert result["box_deficit_px"] == result["target"]["box_deficit_px"]
    assert result["repair_hint"] != "none"


def test_text_fit_keeps_symbol_prefixed_number_and_unit_runs_together():
    tokens = _tokens("-12.5% +3.2% (2025) 5G A/B")
    assert "-12.5%" in tokens
    assert "+3.2%" in tokens
    assert "(2025)" in tokens
    assert "5G" in tokens
    assert "A/B" in tokens


def test_text_fit_applies_chinese_line_break_punctuation_rules():
    tokens = _tokens("强化管控，防患未然（每日提醒）")
    assert "，" not in tokens
    assert all(not token.startswith(tuple("，。！？；：、）》】")) for token in tokens)
    assert all(not token.endswith(tuple("《【〈「（(")) for token in tokens)


def test_text_fit_blocks_shrink_that_destroys_title_hierarchy():
    args = argparse.Namespace(
        text="这是一个必须保持醒目层级的超长主标题", box="150x30",
        font="Microsoft YaHei", font_file=None, bold=True, min_pt=8.0,
        max_pt=32.0, max_lines=1, target_lines=1, scan_step=0.25,
        line_spacing=1.06, width_safety=0.92, height_safety=0.95,
        render_fudge=1.01, target_pt=32.0, semantic_role="hero_title",
        min_readable_pt=None, min_hierarchy_scale=None,
        slide_px="1672x941", slide_in="13.333333x7.505",
    )
    result = best_fit(args)
    assert not result["hierarchy_preserved"]
    assert not result["font_shrink_allowed"]
    assert result["fit_decision"] == "geometry_repair_required"
    assert result["hierarchy_constraints"]["min_readable_pt"] == 24.0


def test_text_fit_preserves_explicit_reference_line_topology():
    args = argparse.Namespace(
        text="第一行\n第二行", box="600x160", font="Microsoft YaHei", font_file=None,
        bold=False, min_pt=8.0, max_pt=26.0, max_lines=2, target_lines=2,
        scan_step=0.25, line_spacing=1.06, width_safety=0.92,
        height_safety=0.95, render_fudge=1.01, target_pt=24.0,
        slide_px="1672x941", slide_in="13.333333x7.505",
    )
    result = best_fit(args)
    assert result["target_lines"] == 2
    assert result["line_topology_preserved"]
    assert result["target"]["line_topology_preserved"]
    assert result["target"]["line_count"] == 2
    assert result["settings"]["measurement_count"] <= 24
    assert result["settings"]["search_strategy"] == "monotonic-boundary-local-refine-v1"


def test_text_fit_accepts_leading_symbol_text_without_cli_rewrite():
    args = argparse.Namespace(
        text="-12.5%", box="220x50", font="Microsoft YaHei", font_file=None,
        bold=False, min_pt=8.0, max_pt=26.0, max_lines=1, target_lines=1,
        scan_step=0.25, line_spacing=1.06, width_safety=0.92,
        height_safety=0.95, render_fudge=1.01, target_pt=20.0,
        slide_px="1672x941", slide_in="13.333333x7.505",
    )
    result = best_fit(args)
    assert result["target"]["lines"] == ["-12.5%"]
    assert result["target"]["fits"]


def test_detected_grid_handles_outer_margins_and_records_expected_centers():
    image = Image.new("RGBA", (330, 210), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    for cy in (45, 160):
        for cx in (45, 155, 285):
            draw.ellipse((cx - 15, cy - 12, cx + 15, cy + 12), fill=(255, 255, 255, 255))
    rows, cols, row_edges, col_edges = _detect_grid(image, 2, 3, 12)
    assert len(rows) == 2 and len(cols) == 3
    assert rows[0] < rows[1] and cols == sorted(cols)
    assert row_edges[0] == 0 and row_edges[-1] == 210
    assert col_edges[0] == 0 and col_edges[-1] == 330


def test_centroid_repack_and_asset_validator_enforce_safe_alpha():
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        source = Image.new("RGBA", (70, 50), (0, 0, 0, 0))
        ImageDraw.Draw(source).polygon([(4, 8), (42, 24), (4, 42)], fill=(255, 255, 255, 255))
        packed, metrics = _square_repack(source, 12, 10)
        path = root / "icon-arrow.png"
        packed.save(path)
        report = validate(root, min_padding=10, max_centroid_delta=0.08)
        assert report["valid"], report
        assert max(abs(value) for value in metrics["centroid_delta_px"]) <= 1.0


def test_alpha_centroid_fit_places_asymmetric_asset_by_visible_subject():
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "asset.png"
        _rgba(path, offset=4)
        x, y, width, height = alpha_centroid_fit(path, 1000, 2000, 4000, 3000)
        assert width > 4000 and height > 3000
        assert (x, y, width, height) != (1000, 2000, 4000, 3000)
        assert y < 2000


def test_native_chart_gap_removes_future_zero_points_from_xml_and_workbook():
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        source = root / "chart-source.pptx"
        repaired = root / "chart-repaired.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        data = ChartData()
        data.categories = [f"M{index}" for index in range(1, 14)]
        data.add_series("2025", tuple(range(120, 107, -1)))
        data.add_series("2026", (96, 91, 87, 82, 78, 73, 69, 0, 0, 0, 0, 0, 0))
        slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(1), Inches(8), Inches(4), data)
        presentation.save(source)

        report = patch_chart(source, repaired, chart_index=0, series_index=1, first_blank_index=7)
        assert report["valid"]
        assert report["patched_series_count"] == 7
        assert report["display_blanks_as"] == "gap"

        with zipfile.ZipFile(repaired, "r") as package:
            chart_xml = package.read(report["chart_path"]).decode("utf-8")
            workbook = load_workbook(io.BytesIO(package.read(report["workbook_path"])), data_only=False)
        assert "$C$2:$C$8" in chart_xml
        assert "$C$2:$C$14" not in chart_xml
        assert 'dispBlanksAs val="gap"' in chart_xml
        sheet = workbook["ChartData"]
        assert [sheet.cell(row, 3).value for row in range(2, 9)] == [96, 91, 87, 82, 78, 73, 69]
        assert all(sheet.cell(row, 3).value is None for row in range(9, 15))


def _deck(path: Path, icon_path: Path, bad_order: bool) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if bad_order:
        icon = slide.shapes.add_picture(str(icon_path), Inches(1), Inches(1), Inches(1), Inches(1))
        icon.name = "icon-main"
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(2))
    card.name = "card-main"
    if not bad_order:
        icon = slide.shapes.add_picture(str(icon_path), Inches(1), Inches(1), Inches(1), Inches(1))
        icon.name = "icon-main"
    text = slide.shapes.add_textbox(Inches(1.5), Inches(1), Inches(1), Inches(0.4))
    text.name = "label-main"
    text.text = "文字"
    presentation.save(path)


def test_pptx_physical_layer_audit_catches_hidden_icon_risk():
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        icon = root / "icon.png"
        _rgba(icon)
        good = root / "good.pptx"
        bad = root / "bad.pptx"
        _deck(good, icon, False)
        _deck(bad, icon, True)
        assert audit(good)["valid"]
        report = audit(bad)
        assert not report["valid"]
        assert "native_object_above_icon" in {item["code"] for item in report["issues"]}


if __name__ == "__main__":
    for test_name, test in sorted(globals().items()):
        if test_name.startswith("test_") and callable(test):
            test()
    print("Knight fidelity port tests passed")
