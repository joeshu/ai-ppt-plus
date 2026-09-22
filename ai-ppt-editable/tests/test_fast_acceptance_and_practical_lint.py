from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"


def test_compact_acceptance_keeps_decisions_and_drops_bulk_details():
    sys.path.insert(0, str(SCRIPTS))
    from compact_acceptance_report import compact
    report = {
        "execution_profile": "fast",
        "phases": {"reference": {"status": "done"}, "final_pptx": True},
        "performance": {"candidate_build_count": 1, "full_render_count": 1, "authoritative_visual_renderer": "powerpoint-export"},
        "benchmark": {"imagegen_call_count": 2, "native_text_coverage": 1.0, "material_mismatch_count": 0},
        "pages": [{"local_crops": [{}, {}, {}]}],
        "hard_blockers": [],
        "final": {"pptx_path": "final.pptx", "render_path": "slide-1.png"},
        "large_debug_payload": ["unused"] * 100,
    }
    result = compact(report)
    assert result["status"] == "passed"
    assert result["evidence"]["local_crop_count"] == 3
    assert "large_debug_payload" not in result
    assert result["full_report_retained"] is True


def test_practical_lint_flags_full_slide_shape_background():
    sys.path.insert(0, str(SCRIPTS))
    from ppt_practical_lint import lint
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "deck.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, deck.slide_height)
        deck.save(path)
        result = lint(path)
        assert "full_slide_shape_background" in {item["code"] for item in result["issues"]}


def test_practical_lint_flags_text_arrow_and_tight_alpha_asset():
    sys.path.insert(0, str(SCRIPTS))
    from ppt_practical_lint import lint
    with tempfile.TemporaryDirectory() as temp_dir:
        temp = Path(temp_dir)
        asset = temp / "tight.png"
        image = Image.new("RGBA", (32, 32), (0, 0, 0, 0))
        ImageDraw.Draw(image).rectangle((0, 8, 20, 24), fill=(255, 0, 0, 255))
        image.save(asset)
        path = temp / "deck.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text = "→"
        slide.shapes.add_picture(str(asset), Inches(2), Inches(2))
        deck.save(path)
        codes = {item["code"] for item in lint(path)["issues"]}
        assert "text_glyph_used_as_arrow" in codes
        assert "transparent_subject_touches_edge" in codes


def test_powerpoint_only_renderer_fails_cleanly_off_windows():
    if sys.platform.startswith("win"):
        return
    with tempfile.TemporaryDirectory() as temp:
        deck = Path(temp) / "empty.pptx"
        Presentation().save(deck)
        result = subprocess.run([
            sys.executable, str(SCRIPTS / "render_authoritative.py"), str(deck),
            "--output-dir", str(Path(temp) / "render"), "--backend", "powerpoint",
        ], capture_output=True, text=True, check=False)
        assert result.returncode == 2
        payload = json.loads(result.stdout)
        assert payload["renderer_policy"] == "powerpoint-first-libreoffice-fallback"
        assert payload["backend_attempts"][0]["reason"] == "not_windows"


def test_compact_acceptance_understands_pipeline_results():
    sys.path.insert(0, str(SCRIPTS))
    from compact_acceptance_report import compact
    passed = compact({"technical_valid": True, "steps": [{"name": "render", "ok": True}]})
    blocked = compact({"technical_valid": False, "steps": [{"name": "render", "ok": False}]})
    assert passed["status"] == "passed" and passed["source_report_kind"] == "pipeline-result"
    assert blocked["status"] == "blocked" and blocked["incomplete_phases"] == ["render"]


if __name__ == "__main__":
    for name, function in sorted(globals().items()):
        if name.startswith("test_") and callable(function):
            function()
    print("fast acceptance and practical lint tests passed")
