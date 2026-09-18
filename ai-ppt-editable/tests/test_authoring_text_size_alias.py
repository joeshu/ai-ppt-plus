import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from scripts.authoring_backend import _normalize_text_size_aliases, build_pptx


class AuthoringTextSizeAliasTest(unittest.TestCase):
    def test_font_size_pt_alias_is_normalized_without_overriding_explicit_size(self):
        slide = {
            "texts": [
                {"object_id": "a", "font_size_pt": 9.5, "text": "A"},
                {"object_id": "b", "font_size_pt": 9.5, "size": 11, "text": "B"},
                {"object_id": "c", "font_size_pt": 10, "runs": [{"text": "C1", "font_size_pt": 8}, {"text": "C2", "size": 12}]},
            ]
        }
        normalized = _normalize_text_size_aliases(slide)
        self.assertEqual(normalized["texts"][0]["size"], 9.5)
        self.assertEqual(normalized["texts"][1]["size"], 11)
        self.assertEqual(normalized["texts"][2]["size"], 10)
        self.assertEqual(normalized["texts"][2]["runs"][0]["size"], 8)
        self.assertEqual(normalized["texts"][2]["runs"][1]["size"], 12)
        self.assertNotIn("size", slide["texts"][0])

    def test_build_pptx_uses_font_size_pt_instead_of_18pt_fallback(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            output = root / "out.pptx"
            deck = {
                "slide_width_in": 13.333333,
                "slide_height_in": 7.5,
                "units": "fraction",
                "assets_dir": str(root),
                "theme": {"font": "Noto Sans CJK SC", "text_color": "#111111"},
                "slides": [{"texts": [{"object_id": "t", "text": "Measured text", "x": 0.1, "y": 0.1, "w": 0.4, "h": 0.1, "font_size_pt": 9.5}]}],
            }
            build_pptx(deck, output)
            prs = Presentation(output)
            shape = next(shape for shape in prs.slides[0].shapes if shape.name == "t")
            run = shape.text_frame.paragraphs[0].runs[0]
            self.assertAlmostEqual(run.font.size.pt, 9.5, places=2)


if __name__ == "__main__":
    unittest.main()
