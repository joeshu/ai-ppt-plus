from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pathlib import Path

src=Path('/mnt/data/latest-knight-contract-test/降套管控_最新Knight流程测试.pptx')
out=Path('/mnt/data/latest-knight-contract-test/降套管控_最新Knight流程测试_r5关键对象修复.pptx')
prs=Presentation(src)
slide=prs.slides[0]

def remove(sh):
    el=sh._element; el.getparent().remove(el)

# Remove wrong large icon pieces, keep the stronger incumbent slogan.
for i,sh in list(enumerate(slide.shapes)):
    if i in {105,106}:
        remove(sh)

# Independent transparent ImageGen asset: corrected calendar + megaphone.
icon='/mnt/data/a_clean_graphic_icon_illustration_on_a_transparent.png'
slide.shapes.add_picture(icon, Emu(10304000), Emu(3130000), width=Emu(1030000), height=Emu(1030000))

# Replace weak checkbox glyphs with editable native red checkboxes.
to_remove=[]
for sh in slide.shapes:
    if hasattr(sh,'text') and sh.text.strip()=='☑':
        to_remove.append(sh)
for sh in to_remove:
    x,y=sh.left,sh.top
    remove(sh)
    box=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x+Emu(90000), y+Emu(100000), Emu(120000), Emu(120000))
    box.fill.background(); box.line.color.rgb=RGBColor(230,0,18); box.line.width=Pt(1.1)
    tick=slide.shapes.add_textbox(x+Emu(85000), y+Emu(50000), Emu(160000), Emu(190000))
    tf=tick.text_frame; tf.clear(); tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p=tf.paragraphs[0]; r=p.add_run(); r.text='✓'; r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=RGBColor(230,0,18)

prs.save(out)
print(out)
