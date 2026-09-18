#!/usr/bin/env python3
"""Place raster/vector assets and manage SVG conversion lifetimes."""
from __future__ import annotations

import os
import math
import posixpath
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from atomic_output import atomic_rewrite_zip
from pptx_primitives import set_alt_text


def _die(message: str, code: int = 2):
    print(f"Error: {message}", file=sys.stderr)
    raise SystemExit(code)


def _resolve(assets_dir: Path, file: str) -> Path:
    path = Path(file)
    return path if path.is_absolute() else assets_dir / path


def _frac(deck: dict, item: dict, key: str, reference: float) -> float:
    if key not in item:
        _die(f"asset is missing coordinate: {key}")
    try:
        value = float(item[key])
    except (TypeError, ValueError):
        _die(f"asset coordinate {key} must be numeric")
    units = deck.get("units", "fraction")
    if units not in {"fraction", "px"}:
        _die(f"unsupported coordinate units: {units}")
    if units == "px":
        if not reference or not math.isfinite(reference):
            _die(f"pixel asset coordinate {key} requires a positive reference canvas")
        value /= reference
    if not math.isfinite(value):
        _die(f"asset coordinate {key} must be finite")
    if key in {"w", "h"} and value <= 0:
        _die(f"asset coordinate {key} must be positive")
    return value


def _validate_box(deck: dict, item: dict, x: float, y: float, w: float, h: float, label: str) -> None:
    if not deck.get("strict_input") or item.get("allow_bleed") is True:
        return
    if x < 0 or y < 0 or x + w > 1 or y + h > 1:
        _die(f"{label} box must stay within the slide in strict input mode: {(x, y, w, h)}")


def replace_svg_media(pptx_path: Path, svg_assets: list[tuple[int, str, Path]]) -> None:
    if not svg_assets:
        return
    presentation_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    relationship_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(pptx_path, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    replacements = []
    for slide_no, object_name, svg_path in svg_assets:
        slide_name = f"ppt/slides/slide{slide_no}.xml"
        relationships_name = f"ppt/slides/_rels/slide{slide_no}.xml.rels"
        root = ET.fromstring(entries[slide_name])
        relationships = ET.fromstring(entries[relationships_name])
        target = None
        for picture in root.findall(f".//{{{presentation_ns}}}pic"):
            non_visual = picture.find(f".//{{{presentation_ns}}}cNvPr")
            if non_visual is None or non_visual.get("name") != object_name:
                continue
            blip = picture.find(f".//{{{drawing_ns}}}blip")
            relationship_id = blip.get(f"{{{relationship_ns}}}embed") if blip is not None else None
            for relationship in relationships:
                if relationship.get("Id") == relationship_id:
                    target = posixpath.normpath(posixpath.join("ppt/slides", relationship.get("Target")))
                    relationship.set("Target", "../media/" + Path(target).with_suffix(".svg").name)
                    break
            break
        if target is None or target not in entries:
            _die(f"could not resolve SVG media relationship for {object_name}")
        new_target = str(Path(target).with_suffix(".svg")).replace("\\", "/")
        entries[relationships_name] = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
        entries[new_target] = svg_path.read_bytes()
        del entries[target]
        replacements.append((target, new_target))
    content_types = ET.fromstring(entries["[Content_Types].xml"])
    for old, new in replacements:
        for override in content_types:
            if override.get("PartName") == "/" + old:
                override.set("PartName", "/" + new)
                override.set("ContentType", "image/svg+xml")
    entries["[Content_Types].xml"] = ET.tostring(content_types, encoding="utf-8", xml_declaration=True)
    atomic_rewrite_zip(pptx_path, entries)


def svg_to_png(svg_path: Path, temporary_files: list[Path] | None = None) -> Path:
    handle, name = tempfile.mkstemp(prefix="ai-ppt-svg-", suffix=".png")
    os.close(handle)
    output = Path(name)
    try:
        from cairosvg import svg2png
        svg2png(url=str(svg_path), write_to=str(output))
        if temporary_files is not None:
            temporary_files.append(output)
        return output
    except ImportError:
        for command in (("inkscape", str(svg_path), "--export-filename", str(output)), ("convert", str(svg_path), str(output))):
            try:
                result = subprocess.run(command, capture_output=True, text=True, timeout=30, check=False)
            except (FileNotFoundError, subprocess.TimeoutExpired):
                continue
            if result.returncode == 0 and output.is_file():
                if temporary_files is not None:
                    temporary_files.append(output)
                return output
        output.unlink(missing_ok=True)
        _die(f"SVG asset requires cairosvg, inkscape, or ImageMagick: {svg_path}")
    except Exception:
        output.unlink(missing_ok=True)
        raise
    return output


def cleanup_temporary_files(paths: list[Path]) -> None:
    for path in paths:
        path.unlink(missing_ok=True)


def alpha_geometry(path: Path, *, threshold: int = 8) -> dict:
    """Measure the visible subject independently from its transparent canvas.

    A small alpha threshold suppresses resampling noise while alpha-weighted
    moments preserve anti-aliased edges.  Returned values are source-pixel
    coordinates and are suitable for deterministic placement/QA evidence.
    """
    try:
        import numpy as np
        from PIL import Image
    except ImportError:
        _die("alpha geometry requires Pillow and numpy")
    with Image.open(path) as source:
        rgba = source.convert("RGBA")
    alpha = np.asarray(rgba.getchannel("A"), dtype=np.float64)
    mask = alpha >= max(1, min(255, int(threshold)))
    ys, xs = np.where(mask)
    if not len(xs):
        _die(f"transparent asset has empty visible alpha: {path}")
    x0, x1 = int(xs.min()), int(xs.max()) + 1
    y0, y1 = int(ys.min()), int(ys.max()) + 1
    weights = np.where(mask, alpha, 0.0)
    total = float(weights.sum())
    grid_y, grid_x = np.indices(alpha.shape)
    centroid_x = float((grid_x * weights).sum() / total)
    centroid_y = float((grid_y * weights).sum() / total)
    return {
        "canvas_px": [rgba.width, rgba.height],
        "visible_bbox_px": [x0, y0, x1 - x0, y1 - y0],
        "visible_centroid_px": [centroid_x, centroid_y],
        "transparent_padding_px": [x0, y0, rgba.width - x1, rgba.height - y1],
        "visible_area_ratio": float(mask.sum() / mask.size),
        "alpha_threshold": int(threshold),
    }


def alpha_centroid_fit(path: Path, x: int, y: int, width: int, height: int, *, threshold: int = 8, contain: float = 1.0) -> tuple[int, int, int, int]:
    """Fit the *visible* alpha bbox into a slot and align visual centroids.

    Unlike canvas-centering, transparent margins do not bias scale or center.
    ``contain`` (<1) reserves symmetric breathing room without changing the
    source asset.  The returned rectangle still places the complete canvas so
    the asset remains independently movable and its alpha is preserved.
    """
    geometry = alpha_geometry(path, threshold=threshold)
    canvas_w, canvas_h = geometry["canvas_px"]
    visible_x, visible_y, visible_w, visible_h = geometry["visible_bbox_px"]
    centroid_x, centroid_y = geometry["visible_centroid_px"]
    contain = max(0.05, min(1.0, float(contain)))
    fit_w, fit_h = width * contain, height * contain
    scale = min(fit_w / visible_w, fit_h / visible_h)
    canvas_out_w = max(1, int(round(canvas_w * scale)))
    canvas_out_h = max(1, int(round(canvas_h * scale)))
    target_cx, target_cy = x + width / 2.0, y + height / 2.0
    placed_x = target_cx - centroid_x * scale
    placed_y = target_cy - centroid_y * scale

    # Centroid alignment is a preference, not permission for the visible
    # subject to escape the declared slot.  Asymmetric alpha mass (for
    # example a skyline above a heavy ribbon) can shift the centroid far from
    # the visible-bbox center.  Clamp the placed canvas just enough to keep
    # the complete visible bbox inside the requested contain box.
    safe_left = x + (width - fit_w) / 2.0
    safe_top = y + (height - fit_h) / 2.0
    safe_right = safe_left + fit_w
    safe_bottom = safe_top + fit_h
    visible_left = placed_x + visible_x * scale
    visible_top = placed_y + visible_y * scale
    visible_right = visible_left + visible_w * scale
    visible_bottom = visible_top + visible_h * scale
    if visible_left < safe_left:
        placed_x += safe_left - visible_left
    elif visible_right > safe_right:
        placed_x -= visible_right - safe_right
    if visible_top < safe_top:
        placed_y += safe_top - visible_top
    elif visible_bottom > safe_bottom:
        placed_y -= visible_bottom - safe_bottom
    return int(round(placed_x)), int(round(placed_y)), canvas_out_w, canvas_out_h


def add_background(slide, slide_spec: dict, assets_dir: Path, sw_emu: int, sh_emu: int):
    from pptx.util import Emu
    background = slide_spec.get("background")
    if not background:
        return
    path = _resolve(assets_dir, background)
    if not path.exists():
        _die(f"background not found: {path}")
    picture = slide.shapes.add_picture(str(path), 0, 0, width=Emu(sw_emu), height=Emu(sh_emu))
    picture.name = str(slide_spec.get("background_object_id", "background"))
    set_alt_text(picture, slide_spec.get("background_alt_text"))


def add_frame(slide, slide_spec: dict, assets_dir: Path, sw_emu: int, sh_emu: int):
    from pptx.util import Emu
    frame = slide_spec.get("frame")
    if not frame:
        return
    path = _resolve(assets_dir, frame)
    if not path.exists():
        _die(f"frame not found: {path}")
    picture = slide.shapes.add_picture(str(path), 0, 0, width=Emu(sw_emu), height=Emu(sh_emu))
    picture.name = str(slide_spec.get("frame_object_id", "frame"))
    set_alt_text(picture, slide_spec.get("frame_alt_text"))


def add_panels(slide, specs: list[dict], assets_dir: Path, deck: dict, ref_w: float, ref_h: float, sw_emu: int, sh_emu: int, slide_no: int):
    from pptx.util import Emu
    for panel in specs:
        path = _resolve(assets_dir, panel["file"])
        if not path.exists():
            _die(f"slide {slide_no}: panel not found: {path}")
        fx, fy = _frac(deck, panel, "x", ref_w), _frac(deck, panel, "y", ref_h)
        fw, fh = _frac(deck, panel, "w", ref_w), _frac(deck, panel, "h", ref_h)
        _validate_box(deck, panel, fx, fy, fw, fh, f"slide {slide_no} panel")
        picture = slide.shapes.add_picture(str(path), Emu(int(fx * sw_emu)), Emu(int(fy * sh_emu)), width=Emu(int(fw * sw_emu)), height=Emu(int(fh * sh_emu)))
        picture.name = str(panel.get("object_id") or panel.get("panel_id") or f"panel-{slide_no}")
        set_alt_text(picture, panel.get("alt_text"))


def add_icons(slide, specs: list[dict], assets_dir: Path, deck: dict, ref_w: float, ref_h: float, sw_emu: int, sh_emu: int, slide_no: int, svg_assets: list[tuple[int, str, Path]], temporary_files: list[Path]):
    from pptx.util import Emu
    for icon_index, icon in enumerate(specs, 1):
        path = _resolve(assets_dir, icon["file"])
        if not path.exists():
            _die(f"slide {slide_no}: icon not found: {path}")
        fx, fy = _frac(deck, icon, "x", ref_w), _frac(deck, icon, "y", ref_h)
        fw, fh = _frac(deck, icon, "w", ref_w), _frac(deck, icon, "h", ref_h)
        _validate_box(deck, icon, fx, fy, fw, fh, f"slide {slide_no} icon")
        source_path = path
        if path.suffix.casefold() == ".svg":
            source_path = svg_to_png(path, temporary_files)
        target = (int(fx * sw_emu), int(fy * sh_emu), int(fw * sw_emu), int(fh * sh_emu))
        if icon.get("placement_mode") == "alpha-centroid-fit" or icon.get("align_by_alpha") is True:
            target = alpha_centroid_fit(source_path, *target, threshold=int(icon.get("alpha_threshold", 8)), contain=float(icon.get("visible_contain", 1.0)))
        picture = slide.shapes.add_picture(str(source_path), Emu(target[0]), Emu(target[1]), width=Emu(target[2]), height=Emu(target[3]))
        picture.name = str(icon.get("name") or icon.get("object_id") or f"icon-{icon_index:02d}")
        set_alt_text(picture, icon.get("alt_text"))
        if path.suffix.casefold() == ".svg":
            svg_assets.append((slide_no, picture.name, path))
