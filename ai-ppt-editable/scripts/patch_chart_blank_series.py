#!/usr/bin/env python3
"""Repair native PPTX chart gaps without materializing missing points as zero."""
from __future__ import annotations

import argparse
import io
import json
import os
import posixpath
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_WORKBOOK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ET.register_namespace("c", C_NS)
ET.register_namespace("r", R_NS)


def _q(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _series_nodes(root: ET.Element) -> list[ET.Element]:
    return root.findall(f".//{_q(C_NS, 'ser')}")


def _cache_points(cache: ET.Element | None) -> list[str | None]:
    if cache is None:
        return []
    count_node = cache.find(_q(C_NS, "ptCount"))
    count = int(count_node.get("val", "0")) if count_node is not None else 0
    values: dict[int, str] = {}
    for point in cache.findall(_q(C_NS, "pt")):
        idx = int(point.get("idx", "0"))
        value = point.find(_q(C_NS, "v"))
        values[idx] = "" if value is None or value.text is None else value.text
        count = max(count, idx + 1)
    return [values.get(index) for index in range(count)]


def _read_axis_values(container: ET.Element | None) -> list[str | None]:
    if container is None:
        return []
    for name in ("strLit", "numLit", "strRef", "numRef"):
        child = container.find(_q(C_NS, name))
        if child is None:
            continue
        cache = child.find(_q(C_NS, "strCache")) or child.find(_q(C_NS, "numCache"))
        if name.endswith("Lit"):
            cache = child
        return _cache_points(cache)
    return []


def _series_name(series: ET.Element, index: int) -> str:
    tx = series.find(_q(C_NS, "tx"))
    if tx is not None:
        value = tx.find(f".//{_q(C_NS, 'v')}")
        if value is not None and value.text:
            return value.text
    return f"Series {index + 1}"


def _excel_col(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _build_cache(parent: ET.Element, values: list[str | None], *, string: bool) -> None:
    cache = ET.SubElement(parent, _q(C_NS, "strCache" if string else "numCache"))
    ET.SubElement(cache, _q(C_NS, "ptCount"), {"val": str(len(values))})
    for index, value in enumerate(values):
        if value in (None, ""):
            continue
        point = ET.SubElement(cache, _q(C_NS, "pt"), {"idx": str(index)})
        ET.SubElement(point, _q(C_NS, "v")).text = str(value)


def _replace_axis_with_ref(container: ET.Element, values: list[str | None], formula: str, *, string: bool) -> None:
    for child in list(container):
        if child.tag in {_q(C_NS, "strLit"), _q(C_NS, "numLit"), _q(C_NS, "strRef"), _q(C_NS, "numRef")}:
            container.remove(child)
    ref = ET.SubElement(container, _q(C_NS, "strRef" if string else "numRef"))
    ET.SubElement(ref, _q(C_NS, "f")).text = formula
    _build_cache(ref, values, string=string)


def _workbook_bytes(categories: list[str | None], series_values: list[list[str | None]], names: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ChartData"
    ws.cell(1, 1, "Category")
    for col, name in enumerate(names, 2):
        ws.cell(1, col, name)
    rows = max([len(categories), *[len(values) for values in series_values]], default=0)
    for row in range(rows):
        ws.cell(row + 2, 1, categories[row] if row < len(categories) else None)
        for col, values in enumerate(series_values, 2):
            value = values[row] if row < len(values) else None
            if value in (None, ""):
                continue
            try:
                ws.cell(row + 2, col, float(value))
            except (TypeError, ValueError):
                ws.cell(row + 2, col, value)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _rels_path(chart_path: str) -> str:
    directory, filename = posixpath.split(chart_path)
    return posixpath.join(directory, "_rels", filename + ".rels")


def _relative_target(source_part: str, target_part: str) -> str:
    return posixpath.relpath(target_part, posixpath.dirname(source_part))


def _parse_or_new_rels(raw: bytes | None) -> ET.Element:
    return ET.fromstring(raw) if raw else ET.Element(_q(PKG_REL_NS, "Relationships"))


def _next_rel_id(root: ET.Element) -> str:
    used = {child.get("Id") for child in root}
    number = 1
    while f"rId{number}" in used:
        number += 1
    return f"rId{number}"


def _ensure_xlsx_content_type(raw: bytes) -> bytes:
    ET.register_namespace("", CT_NS)
    root = ET.fromstring(raw)
    if not any(child.tag == _q(CT_NS, "Default") and child.get("Extension") == "xlsx" for child in root):
        ET.SubElement(root, _q(CT_NS, "Default"), {"Extension": "xlsx", "ContentType": XLSX_CT})
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def patch_chart(pptx_in: Path, pptx_out: Path, *, chart_index: int, series_index: int, first_blank_index: int) -> dict:
    with zipfile.ZipFile(pptx_in, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    chart_paths = sorted(name for name in entries if name.endswith(".xml") and "/chart" in name and "/_rels/" not in name and posixpath.basename(name).startswith("chart"))
    if not chart_paths:
        raise ValueError("no native chart XML parts found")
    if not 0 <= chart_index < len(chart_paths):
        raise IndexError(f"chart_index {chart_index} outside 0..{len(chart_paths)-1}")
    chart_path = chart_paths[chart_index]
    root = ET.fromstring(entries[chart_path])
    series = _series_nodes(root)
    if not 0 <= series_index < len(series):
        raise IndexError(f"series_index {series_index} outside 0..{len(series)-1}")
    categories = _read_axis_values(series[0].find(_q(C_NS, "cat")))
    values = [_read_axis_values(node.find(_q(C_NS, "val"))) for node in series]
    names = [_series_name(node, index) for index, node in enumerate(series)]
    if first_blank_index < 1 or first_blank_index > len(categories):
        raise ValueError("first_blank_index must be between 1 and category count")
    values[series_index] = values[series_index][:first_blank_index]
    workbook_path = f"ppt/embeddings/chart-data-snapshot-{chart_index + 1:03d}.xlsx"
    entries[workbook_path] = _workbook_bytes(categories, values, names)
    cat_formula = f"'ChartData'!$A$2:$A${len(categories)+1}"
    for index, node in enumerate(series):
        cat = node.find(_q(C_NS, "cat"))
        val = node.find(_q(C_NS, "val"))
        if cat is not None:
            _replace_axis_with_ref(cat, categories, cat_formula, string=True)
        if val is not None:
            column = _excel_col(index + 2)
            _replace_axis_with_ref(val, values[index], f"'ChartData'!${column}$2:${column}${len(values[index])+1}", string=False)
    rels_path = _rels_path(chart_path)
    rels = _parse_or_new_rels(entries.get(rels_path))
    external = root.find(_q(C_NS, "externalData"))
    if external is None:
        external = ET.SubElement(root, _q(C_NS, "externalData"))
    rel_id = external.get(_q(R_NS, "id"))
    if not rel_id:
        rel_id = _next_rel_id(rels)
        external.set(_q(R_NS, "id"), rel_id)
    matching = next((relation for relation in rels if relation.get("Id") == rel_id), None)
    if matching is None:
        matching = ET.SubElement(rels, _q(PKG_REL_NS, "Relationship"), {"Id": rel_id})
    matching.set("Type", REL_WORKBOOK)
    matching.set("Target", _relative_target(chart_path, workbook_path))
    auto = external.find(_q(C_NS, "autoUpdate"))
    if auto is None:
        auto = ET.SubElement(external, _q(C_NS, "autoUpdate"))
    auto.set("val", "0")
    blank = root.find(_q(C_NS, "dispBlanksAs"))
    if blank is None:
        blank = ET.SubElement(root, _q(C_NS, "dispBlanksAs"))
    blank.set("val", "gap")
    entries[chart_path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    entries[rels_path] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    entries["[Content_Types].xml"] = _ensure_xlsx_content_type(entries["[Content_Types].xml"])
    pptx_out.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(prefix=pptx_out.name + ".", suffix=".tmp", dir=pptx_out.parent)
    os.close(handle)
    try:
        with zipfile.ZipFile(temp_name, "w", zipfile.ZIP_DEFLATED) as target:
            for name, payload in entries.items():
                target.writestr(name, payload)
        os.replace(temp_name, pptx_out)
    finally:
        Path(temp_name).unlink(missing_ok=True)
    return {"schema": "ai-ppt-plus/chart-blank-series/v1", "valid": True, "input": str(pptx_in), "output": str(pptx_out), "chart_path": chart_path, "chart_index": chart_index, "series_index": series_index, "series_name": names[series_index], "category_count": len(categories), "patched_series_count": len(values[series_index]), "first_blank_index": first_blank_index, "display_blanks_as": "gap", "workbook_path": workbook_path}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--chart-index", type=int, default=0)
    parser.add_argument("--series-index", type=int, required=True)
    parser.add_argument("--first-blank-index", type=int, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    try:
        report = patch_chart(args.input.resolve(), args.output.resolve(), chart_index=args.chart_index, series_index=args.series_index, first_blank_index=args.first_blank_index)
    except Exception as exc:
        print(json.dumps({"schema": "ai-ppt-plus/chart-blank-series/v1", "valid": False, "status": "blocked", "code": "chart_blank_series_patch_failed", "message": str(exc)}, ensure_ascii=False))
        return 2
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
